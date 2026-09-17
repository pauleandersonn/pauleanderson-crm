# -*- coding: utf-8 -*-
"""
Gera proposta-MCR-Academy-v1.docx
Alteracoes vs versao anterior (HTML 07/07/2026):
  - 4h por dia, 1 encontro/semana, 10 semanas (40h mantidas)
  - Turma aberta (sem numero fixo de alunas)
  - Divisao de turma "a verificar"
  - Identidade: MCR Academy - Escola de Comunicacao e Midia
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\paule\Documents\PROGRAMAÇÃO\PROPOSTAS\INSTITUTO-CONSULADO-MULHER\proposta-MCR-Academy-v1.docx"

doc = Document()

# Margens
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# Estilo base
style = doc.styles['Normal']
style.font.name = 'Georgia'
style.font.size = Pt(11)

def heading(text, level=1, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(6)
    if color:
        run.font.color.rgb = color
    return p

def para(text, italic=False, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    return p

def bullet(text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    return p

# === CAPA / HERO ===
heading('MCR ACADEMY', level=1, color=RGBColor(0xA0, 0x7F, 0x3A))
heading('Escola de Comunicação e Mídia', level=2)
p = doc.add_paragraph()
r = p.add_run('Proposta Comercial · 08 de julho de 2026')
r.italic = True
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

heading('Formação Fotografia da Mulher Empreendedora', level=1)

p = doc.add_paragraph()
r = p.add_run(
    'Programa de 40 horas em 10 encontros práticos semanais, do celular ao portfólio '
    '— pensado para mulheres em situação de vulnerabilidade que já empreendem ou querem empreender.'
)
r.font.size = Pt(11)

p = doc.add_paragraph()
for pill in ['Carga: 40h', '10 encontros', '1 encontro/semana', '4h por encontro', '2026']:
    r = p.add_run(f'  {pill}  ')
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xA0, 0x7F, 0x3A)

doc.add_paragraph()

# === 1. APRESENTAÇÃO ===
heading('1. Apresentação', level=1)
p = doc.add_paragraph()
r = p.add_run(
    'O celular é o único estúdio que elas já têm. A formação existe pra transformar '
    'esse estúdio em ferramenta de trabalho.'
)
r.italic = True

heading('Base de experiência comprovada', level=3)
para('Ao longo da minha trajetória, já passaram pelas minhas aulas mais de 1.000 alunos — '
     'entre turmas de formação inicial, oficinas livres e cursos institucionalizados. '
     'Duas instituições de destaque:')
bullet('CETAM (Centro de Educação Tecnológica do Amazonas) — instrutor de fotografia '
       'aplicada, turmas com mais de 60 alunos cada, em 4 edições consecutivas. '
       'Carga horária total acumulada acima de 240h só nessa frente.')
bullet('CIEEI (Centro Integrado de Educação e Esporte da Infância) — oficinas de '
       'comunicação visual e fotografia pra jovens em formação profissional.')
para('Não é palestra teórica. É prática testada com públicos grandes, em salas reais, '
     'com gente que nunca teve aula de imagem antes.')

# === 1.5 PORTFÓLIO ===
heading('1.5 Portfólio · Onde você pode ver meu trabalho', level=1)
para('Antes de ler o resto, dá uma olhada no que já produzi. Tudo público, tudo verificável:')
bullet('YouTube: youtube.com/@Pauleandersonsouza — vídeos de projetos, aulas e bastidores')
bullet('Site institucional: pauloanderson.plano.wordpress.com — portfólio completo, serviços e cases')
bullet('Instagram pessoal: @pauleandersonn — dia a dia, bastidores, ensaios')
bullet('Instagram de filmes: @pauloandersonfilmes — conteúdo audiovisual, projetos de vídeo')
para('Se preferir conversar antes, meu WhatsApp é (92) 99241-1099.')

# === 2. ENTENDIMENTO ===
heading('2. Entendimento do desafio', level=1)
para('Esta proposta responde a três perguntas concretas:')
bullet('Como ensinar fotografia com celular a alguém que nunca teve aula de imagem?')
bullet('Como transformar essa habilidade em ativo de comunicação para um negócio real '
       '(mesmo que pequeno)?')
bullet('Como produzir um portfólio final que essas mulheres possam usar imediatamente '
       '— em Instagram, em proposta pra cliente, em banca de microcrédito?')

# === 3. METODOLOGIA ===
heading('3. Metodologia · Três pilares', level=1)
para('A formação segue o método que desenvolvi — Inspirar · Comunicar · Transformar — '
     'adaptado ao contexto do Instituto Consulado da Mulher (linguagem laica, '
     'foco em autonomia econômica, sem nenhuma dimensão confessional).')

t = doc.add_table(rows=4, cols=3)
t.style = 'Light Grid Accent 1'
hdr = t.rows[0].cells
hdr[0].text = 'Pilar'
hdr[1].text = 'O que a aluna ganha'
hdr[2].text = 'Como aparece no programa'
for cell in hdr:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

rows = [
    ('Inspirar', 'Confiança mínima no próprio olhar',
     'Encontro 1 — "Aprender a ver antes de fotografar"'),
    ('Comunicar', 'Técnica de celular aplicável',
     'Encontros 2 a 9 — iluminação, cenário, composição, edição, narrativa'),
    ('Transformar', 'Portfólio real pronto pra uso',
     'Encontro 10 — fechamento com 5 fotos curadas por aluna'),
]
for i, (a, b, c) in enumerate(rows, 1):
    cells = t.rows[i].cells
    cells[0].text = a
    cells[1].text = b
    cells[2].text = c

# === 4. PROGRAMAÇÃO ===
heading('4. Programação completa — 40 horas, 10 encontros', level=1)
para('Formato: 1 encontro por semana, 4 horas por encontro. Pode ser ajustado de acordo '
     'com a demanda do Instituto. Todos os encontros têm conteúdo expositivo/demonstrativo '
     '+ prática assistida + roda de devolutiva. Pensado pra quem nunca fotografou com intenção.')

t = doc.add_table(rows=11, cols=5)
t.style = 'Light Grid Accent 1'
hdr = t.rows[0].cells
hdr[0].text = '#'
hdr[1].text = 'Encontro'
hdr[2].text = 'Carga'
hdr[3].text = 'Conteúdo'
hdr[4].text = 'Entregável'
for cell in hdr:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

encontros = [
    ('1', 'Aprender a ver', '4h',
     'Leitura de imagem, regra dos terços, olhar atento, apresentação do projeto pessoal',
     '10 fotos de exercício (tema "minha rotina")'),
    ('2', 'Luz natural e iluminação com celular', '4h',
     'Horários de ouro, luz dura vs. suave, sombras como recurso, app de lanterna doméstica',
     '10 fotos de exercício (tema "luz da minha casa")'),
    ('3', 'Cenário e composição', '4h',
     'Fundos limpos, regra dos terços avançada, simetria, enquadramento pra Instagram (4:5, 1:1)',
     '10 fotos de um produto/hobby'),
    ('4', 'Retrato e autorretrato', '4h',
     'Enquadramento de pessoa, expressão, fundo desfocado (modo retrato), celular na mão da outra',
     '10 fotos (retrato de alguém querido)'),
    ('5', 'Edição leve no celular', '4h',
     'Snapseed (grátis): ajuste de luz, corte, nitidez. App nativo de fotos. Truques para fotos escuras',
     'Antes/depois lado a lado'),
    ('6', 'Contar uma história em 5 fotos', '4h',
     'Narrativa visual sequencial, legenda que conecta, uso prático (Instagram, WhatsApp Status, proposta comercial)',
     'Mini-ensaio de 5 fotos com legenda'),
    ('7', 'Vender com imagem', '4h',
     'Fotografia de produto pra venda online (WhatsApp, Instagram, marketplace)',
     '10 fotos de produto com legenda de venda'),
    ('8', 'Identidade visual no feed', '4h',
     'Paleta de cor pessoal, consistência, grade do Instagram, capa de destaque',
     'Mockup de feed com 9 fotos'),
    ('9', 'Retrato profissional com celular', '4h',
     'Auto-retrato, retrato de equipe, fundo neutro, luz de janela',
     '5 retratos prontos pra uso'),
    ('10', 'Fechamento · Portfólio', '4h',
     'Curadoria, apresentação oral, plano de continuidade, grupo de WhatsApp de continuidade',
     '5 fotos curadas + 1 página de portfólio impressa'),
]
for i, (n, e, c, ct, en) in enumerate(encontros, 1):
    cells = t.rows[i].cells
    cells[0].text = n
    cells[1].text = e
    cells[2].text = c
    cells[3].text = ct
    cells[4].text = en

heading('Materiais inclusos na formação', level=3)
bullet('Apostila ilustrada (PDF, 24 páginas) com linguagem acessível — enviada por WhatsApp no Encontro 1')
bullet('Lista de 1 app gratuito recomendado por módulo (Snapseed, Lightroom Mobile gratuito, Instagram)')
bullet('Modelo de 1 página de portfólio editável (Canva, gratuito) entregue no Encontro 10')
bullet('Grupo de WhatsApp de continuidade (operação por 60 dias após o Encontro 10, sem custo extra)')
bullet('1 página impressa do portfólio final por aluna (produzida previamente pelo instrutor)')

# === 5. LOGÍSTICA ===
heading('5. Logística e cronograma', level=1)
p = doc.add_paragraph()
r = p.add_run('Formato semanal: ')
r.bold = True
p.add_run('1 encontro por semana · 4 horas por encontro · 10 semanas.')

p = doc.add_paragraph()
r = p.add_run('Ajustável: ')
r.bold = True
p.add_run('a carga horária e o cronograma podem ser ajustados de acordo com a demanda do Instituto.')

p = doc.add_paragraph()
r = p.add_run('Locais propostos (a confirmar): ')
r.bold = True
p.add_run('Espaço Monte das Oliveiras e Santos Dumont — fornecidos pelo Instituto.')

p = doc.add_paragraph()
r = p.add_run('Tamanho da turma: ')
r.bold = True
p.add_run('aberto. O programa se adapta ao número de inscritas definido pelo Instituto.')

p = doc.add_paragraph()
r = p.add_run('Divisão de turmas: ')
r.bold = True
p.add_run('a verificar — definir conforme número final de inscritas, espaço físico disponível e logística do Instituto.')

# === 6. SEGURANÇA ===
heading('6. Segurança e idoneidade · Como você se protege nessa parceria', level=1)
para('Trabalhar com recursos de uma OSCIP exige garantia dos dois lados. Listo abaixo o que '
     'está coberto por padrão nesta proposta — sem custo extra, porque é o mínimo que se '
     'espera de uma parceria institucional séria:')
bullet('Contrato formal de prestação de serviço (1 a 2 páginas), com cronograma, '
       'responsabilidades, forma de pagamento e cláusula de rescisão — assinado pelas '
       'duas partes antes do Encontro 1.')
bullet('Emissão de nota fiscal (NF-e de serviço) pela MCR Academy, com CNPJ ativo. '
       'O Instituto recebe documento fiscal válido pra prestação de contas.')
bullet('Garantia de cumprimento: se eu não entregar o programa conforme cronograma acordado, '
       'o Instituto retém proporcionalmente o último pagamento. Isso está no contrato.')
bullet('Seguro de responsabilidade civil do instrutor: a confirmar no momento da assinatura, '
       'conforme exigência do Instituto (se for exigido, incluyo no custo).')
bullet('Relatório final de prestação de contas em PDF ao final do Encontro 10, com: '
       'lista de presença por encontro, fotos autorizadas, portfólios impressos entregues, '
       'horas realizadas, saldo de pagamentos.')
bullet('Referências verificáveis: posso fornecer contato do CETAM e de outros clientes '
       'institucionais sob solicitação, antes da assinatura do contrato.')

# === 7. INVESTIMENTO ===
heading('7. Investimento', level=1)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Valor único · programa completo')
r.italic = True
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0xA0, 0x7F, 0x3A)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('R$ 6.000')
r.bold = True
r.font.size = Pt(32)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Pagamento total pela formação em 40 horas. Sem cobrança adicional por aluna, '
              'sem taxa de deslocamento dentro de Manaus.')
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

heading('O que está incluso nos R$ 6.000', level=3)
bullet('40 horas de instrução presencial (10 encontros de 4h)')
bullet('12 horas de preparação de material (apostila, exercícios, modelo de portfólio)')
bullet('Apostila ilustrada em PDF (24 páginas) — enviada por WhatsApp')
bullet('1 página de portfólio impressa por aluna (papel couchê)')
bullet('Grupo de WhatsApp de continuidade por 60 dias após o Encontro 10')
bullet('Relatório final de prestação de contas em PDF')
bullet('Contrato formal + NF-e + termo de LGPD')

heading('O que não está incluso (e pode ser contratado à parte, se desejado)', level=3)
bullet('Gravação em vídeo das aulas (orçamento separado, sob demanda)')
bullet('Deslocamento pra cidades fora de Manaus')
bullet('Certificados em papel (o Instituto emite os seus; se preferir que eu produza, é orçamento à parte)')
bullet('Microsites individuais por aluna (formato premium, orçamento à parte)')

heading('Forma de pagamento sugerida', level=3)
t = doc.add_table(rows=4, cols=4)
t.style = 'Light Grid Accent 1'
hdr = t.rows[0].cells
hdr[0].text = 'Parcela'
hdr[1].text = 'Valor'
hdr[2].text = 'Vencimento'
hdr[3].text = 'Condição'
for cell in hdr:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

pagamentos = [
    ('1ª', 'R$ 2.000', 'Na assinatura do contrato', 'Reserva a agenda e inicia preparação'),
    ('2ª', 'R$ 2.000', 'Encontro 5 (meio do programa)', 'Confirma continuidade'),
    ('3ª', 'R$ 2.000', 'Encontro 10 (entrega final)', 'Liberada após a entrega dos portfólios'),
]
for i, (a, b, c, d) in enumerate(pagamentos, 1):
    cells = t.rows[i].cells
    cells[0].text = a
    cells[1].text = b
    cells[2].text = c
    cells[3].text = d

para('Outras condições (à vista com 5% de desconto, ou parcelamento diferente) podem ser '
     'combinadas na conversa de alinhamento.')

# === 8. PARCERIA CONTÍNUA ===
heading('8. Modelo de parceria contínua · 1 projeto por trimestre', level=1)
para('A proposta de parceria anual segue o mesmo modelo de R$ 6.000 por trimestre, com '
     '4 trilhas complementares ao longo de 12 meses. Detalhes sob solicitação.')

# === 9. DIFERENCIAIS ===
heading('9. Diferenciais desta proposta', level=1)
bullet('Experiência prévia comprovada como instrutor pelo CETAM em formação pra mulheres '
       '— não é palestra teórica, é prática testada.')
bullet('Currículo replicável — o material vira módulo de curso online no MCR Academy, '
       'ampliando o impacto social sem custo adicional.')
bullet('Tom respeitoso com o público — método adaptado pra mulheres em vulnerabilidade '
       '(linguagem acessível, sem promessas de "faturar X mil"), com acolhimento de '
       'trajetórias diversas.')
bullet('Entregável concreto, não horas vagas — cada aluna sai com 5 fotos curadas + '
       '1 página de portfólio impressa e pronta pra uso.')
bullet('Idoneidade documental — NF-e emitida pela MCR Academy (CNPJ ativo), contrato '
       'formal e relatório final de prestação de contas. Tudo pronto pra auditoria.')

# === 10. PRÓXIMOS PASSOS ===
heading('10. Próximos passos (sugestão)', level=1)
para('1. Conversa de alinhamento (30-45min, online ou presencial) — pra eu entender: '
     'número de inscritas, datas preferidas, infraestrutura disponível nos locais, '
     'expectativas de portfólio, e tirar dúvidas.')
para('2. Acordo de investimento — valor e forma de pagamento.')
para('3. Assinatura de contrato simples (1 página) com cronograma, responsabilidades e '
     'cláusula de imagem (LGPD).')
para('4. Encontro 1 — aproximadamente 30 dias após a assinatura, com tempo hábil pra '
     'preparar material + organização das turmas.')

# === ASSINATURA ===
doc.add_paragraph()
para('À disposição pra qualquer conversa. Pode me chamar no WhatsApp (92) 99241-1099 ou '
     'responder por aqui.', italic=True)

p = doc.add_paragraph()
r = p.add_run('Pauleanderson Souza')
r.bold = True

para('Publicitário · Fotógrafo · Instrutor')
para('Fundador — MCR Academy — Escola de Comunicação e Mídia')
para('Manaus, AM · Brasil · WhatsApp (92) 99241-1099')

# === RODAPÉ ===
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('MCR ACADEMY')
r.bold = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0xA0, 0x7F, 0x3A)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Escola de Comunicação e Mídia · Inspirar · Comunicar · Transformar')
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Proposta emitida em 08/07/2026 · válida até 08/08/2026')
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.save(OUT)
print(f'OK: {OUT}')